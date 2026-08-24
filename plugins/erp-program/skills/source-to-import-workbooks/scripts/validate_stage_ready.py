#!/usr/bin/env python3
"""
Validate a stage-ready import file against the Import Review contract BEFORE
it is uploaded to /imports. Catching contract violations here is the whole
point: a staged batch with avoidable validation errors wastes a reviewer's
time and erodes trust in the pipeline.

Usage:
    python validate_stage_ready.py <file.xlsx|file.csv> [--family legacy_system]
    python validate_stage_ready.py stage-ready/*.xlsx

Exit code 0 = clean, 1 = problems found (all problems are printed).
The contract is read from ../references/import-contract.json.
"""
import sys, os, json, re, csv, argparse

HERE = os.path.dirname(os.path.abspath(__file__))
CONTRACT_PATH = os.path.join(HERE, "..", "references", "import-contract.json")


def load_contract():
    with open(CONTRACT_PATH, "r", encoding="utf-8") as fh:
        return json.load(fh)


def read_rows(path):
    """Return (headers, rows, warnings). Only the first worksheet is read,
    mirroring Import Review behavior."""
    warnings = []
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as fh:
            r = list(csv.reader(fh))
        if not r:
            return [], [], ["file is empty"]
        return [c.strip() for c in r[0]], r[1:], warnings
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            print("openpyxl is required: pip install openpyxl --break-system-packages")
            sys.exit(2)
        # data_only=False so formulas are detectable
        wb = load_workbook(path, data_only=False)
        if len(wb.worksheets) > 1:
            warnings.append(
                f"workbook has {len(wb.worksheets)} worksheets; Import Review reads only "
                f"the first ('{wb.worksheets[0].title}'). Remove extra sheets from stage-ready files."
            )
        ws = wb.worksheets[0]
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        # formula / merged-cell checks
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    warnings.append(f"formula found at {c.coordinate}; stage-ready files must contain static values")
                    break
        if ws.merged_cells.ranges:
            warnings.append(f"merged cells found ({len(ws.merged_cells.ranges)} range(s)); not allowed in stage-ready files")
        if not grid:
            return [], [], warnings + ["worksheet is empty"]
        headers = ["" if v is None else str(v).strip() for v in grid[0]]
        rows = [["" if v is None else str(v).strip() for v in r] for r in grid[1:]]
        return headers, rows, warnings
    return [], [], [f"unsupported extension '{ext}'; use .csv or .xlsx"]


def detect_family(headers, contract):
    """Guess the family from which required field keys are present."""
    best, best_score = None, 0
    for name, fam in contract["families"].items():
        req = [f for f, spec in fam["fields"].items() if spec.get("required")]
        if not req:
            continue
        score = sum(1 for f in req if f in headers) / len(req)
        if score > best_score:
            best, best_score = name, score
    return (best, best_score) if best_score >= 0.6 else (None, best_score)


def validate(path, contract, family_arg=None):
    problems, notes = [], []
    limits = contract["limits"]

    ext = os.path.splitext(path)[1].lower()
    if ext not in limits["allowedExtensions"]:
        problems.append(f"extension '{ext}' not allowed; use {' or '.join(limits['allowedExtensions'])}")

    size = os.path.getsize(path)
    if size > limits["maxFileBytes"]:
        problems.append(f"file is {size/1048576:.1f} MB; limit is {limits['maxFileBytes']/1048576:.0f} MB")

    headers, rows, warns = read_rows(path)
    problems.extend(warns)
    if not headers:
        return problems, notes
    if not rows:
        problems.append("no data rows found below the header row")
        return problems, notes

    family = family_arg
    if not family:
        family, score = detect_family(headers, contract)
        if not family:
            problems.append(
                "could not determine the record family from the headers "
                f"(best match {score:.0%}). Pass --family explicitly, and check that headers "
                "are exact field keys rather than display labels."
            )
            return problems, notes
        notes.append(f"detected family: {family}")

    if family not in contract["families"]:
        problems.append(f"unknown family '{family}'")
        return problems, notes

    fam = contract["families"][family]
    spec = dict(contract["universalFields"])
    spec.update(fam["fields"])

    # Unknown headers: not fatal (they are dropped on create) but worth surfacing.
    unknown = [h for h in headers if h and h not in spec]
    if unknown:
        notes.append(
            "columns not in the contract (they will NOT be copied onto the created record): "
            + ", ".join(unknown)
        )

    idx = {h: i for i, h in enumerate(headers) if h}
    missing_required = [f for f, s in spec.items() if s.get("required") and f not in idx]
    if missing_required:
        problems.append("missing required column(s): " + ", ".join(sorted(missing_required)))

    fmts = contract["formats"]
    seen_titles = {}

    def cell(row, field):
        i = idx.get(field)
        if i is None or i >= len(row):
            return ""
        return (row[i] or "").strip()

    for n, row in enumerate(rows, start=2):  # row 1 is the header
        if not any((c or "").strip() for c in row):
            continue
        for field, s in spec.items():
            if field not in idx:
                continue
            val = cell(row, field)
            if not val:
                if s.get("required"):
                    problems.append(f"row {n}: '{field}' is required but blank")
                continue
            if "enum" in s and val not in s["enum"]:
                problems.append(
                    f"row {n}: '{field}' = '{val}' is not allowed; use one of: {', '.join(s['enum'])}"
                )
            if val in (s.get("rejectedValues") or []):
                problems.append(
                    f"row {n}: '{field}' = '{val}' is rejected by staging; move this record to the review pack"
                )
            fmt = s.get("format")
            if fmt == "date" and not re.match(fmts["date"]["pattern"], val):
                problems.append(f"row {n}: '{field}' = '{val}' must be ISO YYYY-MM-DD")
            if fmt == "boolean" and val.lower() not in fmts["boolean"]["allowed"]:
                problems.append(
                    f"row {n}: '{field}' = '{val}' must be one of {', '.join(fmts['boolean']['allowed'])}"
                )
            if fmt == "uuid" and not re.match(fmts["uuid"]["pattern"], val):
                problems.append(
                    f"row {n}: '{field}' = '{val}' is not a UUID. workstreamId must be a real workstream "
                    "UUID or left blank — never a workstream name."
                )
        t = cell(row, "title")
        if t:
            if t in seen_titles:
                problems.append(f"row {n}: duplicate title '{t}' (also row {seen_titles[t]}); titles must be unique")
            else:
                seen_titles[t] = n

    notes.append(f"{len([r for r in rows if any((c or '').strip() for c in r)])} data row(s) checked against family '{family}'")
    return problems, notes


def main():
    ap = argparse.ArgumentParser(description="Validate stage-ready import files against the Import Review contract.")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--family", default=None, help="record family; auto-detected when omitted")
    args = ap.parse_args()

    contract = load_contract()
    failed = False
    for path in args.files:
        print(f"\n=== {os.path.basename(path)} ===")
        if not os.path.exists(path):
            print("  ERROR: file not found")
            failed = True
            continue
        problems, notes = validate(path, contract, args.family)
        for nte in notes:
            print(f"  note: {nte}")
        if problems:
            failed = True
            print(f"  {len(problems)} problem(s):")
            for p in problems:
                print(f"    - {p}")
        else:
            print("  OK — contract-clean. Safe to stage in Import Review.")
    print()
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
